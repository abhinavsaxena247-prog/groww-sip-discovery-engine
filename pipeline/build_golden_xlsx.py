"""Golden set workbook: user labels blind; AI labels in a hidden sheet; agreement by formula."""
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

G = json.load(open("/home/claude/engine/out/golden_selection.json"))
GROWW = "027F5F"   # darker Groww green (white text 5.0:1); brand primary #04B488 from groww.in
F = "Arial"
hdr = Font(name=F, bold=True, color="FFFFFF"); body = Font(name=F, size=10)
fill_h = PatternFill("solid", fgColor=GROWW); fill_in = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

EVENTS = [("None of these", "none"), ("Started a SIP", "started"), ("Continued / still investing", "continued"),
          ("Checked or tracked", "checked_tracked"), ("Skipped an instalment", "skipped"), ("Paused SIP", "paused"),
          ("Reduced amount", "reduced"), ("Increased / step-up", "increased_stepup"),
          ("Stopped or cancelled SIP", "stopped_cancelled"), ("Payment failed / SIP lapsed", "lapsed_failed_payment"),
          ("Redeemed / withdrew money", "redeemed_withdrew"), ("Switched platform", "switched_platform")]
DROPS = [("None", "none"), ("Voluntary (they chose to stop/skip/pause/reduce)", "voluntary"),
         ("Involuntary (failed payment/mandate/app error, not their choice)", "involuntary"),
         ("Passive drift (forgot, lost track)", "passive_drift")]
DRIVERS = [("No clear reason", "none"), ("Cash-flow or income change", "cashflow_or_income_change"),
           ("Needed money (emergency/goal)", "emergency_or_goal_need_for_money"),
           ("Returns lower than expected", "returns_below_expectation"),
           ("Negative or volatile returns", "negative_or_volatile_returns"),
           ("Payment / mandate / autopay failure", "payment_mandate_failure"),
           ("App bug or friction", "app_friction_or_bug"),
           ("Missing control (pause/skip/modify)", "missing_control_pause_skip_modify"),
           ("Confusing numbers or terms", "confusing_numbers_or_terms"), ("Fees or hidden charges", "fees_or_hidden_charges"),
           ("Trust or support failure", "trust_or_support_failure"), ("Advice from others / media", "advice_from_others_or_media"),
           ("Goal clarity or tracking", "goal_clarity_or_tracking"), ("Habit / automation (positive)", "habit_or_automation_positive"),
           ("Other", "other")]
FUND = re.compile(r"\b(parag parikh|ppfas|nippon|hdfc|icici|sbi|axis|mirae|quant|kotak|uti|motilal|tata|nifty ?50|small ?cap|mid ?cap|flexi ?cap)\b", re.I)
ADVICE = re.compile(r"(you should (buy|invest|stop|continue)|invest in |switch to |best fund|recommend(ed)? fund|don't stop|keep investing)", re.I)

wb = Workbook()
# ---- Instructions ----
ins = wb.active; ins.title = "Instructions"
lines = [("Golden set: hand labels for the Discovery Engine", True),
         ("What: 26 real public comments. You label each one yourself, without looking at the AI's answer.", False),
         ("Why: your labels are the 'truth'. The sheet then shows how often the AI agreed with you (needed for the Evidence Log).", False),
         ("How: go to the 'Label here' tab. Fill ONLY the 4 yellow columns using the dropdowns. ~20 minutes.", False),
         ("Rules: judge only from the comment text. If unsure, pick the closest option and write why in Notes.", False),
         ("Do NOT open the hidden 'AI labels' tab until you finish. The 'Agreement' tab updates as you fill.", False),
         ("", False), ("Column definitions", True),
         ("1. Relevant? Yes if the writer describes their OWN experience/question/need about SIPs or mutual funds. Generic praise ('good app') or stock trading = No.", False),
         ("2. Main thing that happened to the SIP: the most important event the writer says happened to them.", False),
         ("3. Drop-off type: Voluntary = they chose; Involuntary = it failed/lapsed without their choice; Passive = forgot/lost track; None = no drop-off.", False),
         ("4. Main reason: the main reason the writer gives for what happened. 'No clear reason' if none.", False),
         ("", False), ("Example row (not part of the set)", True),
         ("Text: 'My autopay failed twice this month and now Groww says my SIP is cancelled. I never stopped it!'", False),
         ("→ Relevant = Yes | Event = Payment failed / SIP lapsed | Drop-off = Involuntary | Reason = Payment / mandate / autopay failure", False)]
for i, (t, b) in enumerate(lines, 1):
    c = ins.cell(row=i, column=1, value=t); c.font = Font(name=F, bold=b, size=12 if b else 10, color=GROWW if b else "000000")
    c.alignment = Alignment(wrap_text=True)
ins.column_dimensions["A"].width = 130

# ---- Lists (dropdown sources + key maps) ----
L = wb.create_sheet("Lists")
for col, data in ((1, EVENTS), (3, DROPS), (5, DRIVERS)):
    L.cell(row=1, column=col, value="Label").font = Font(name=F, bold=True)
    L.cell(row=1, column=col + 1, value="Key").font = Font(name=F, bold=True)
    for r, (lab, key) in enumerate(data, 2):
        L.cell(row=r, column=col, value=lab).font = body; L.cell(row=r, column=col + 1, value=key).font = body
L.cell(row=1, column=8, value="YesNo").font = Font(name=F, bold=True)
L.cell(row=2, column=8, value="Yes"); L.cell(row=3, column=8, value="No")
L.sheet_state = "hidden"

# ---- Label here ----
S = wb.create_sheet("Label here", 1)
cols = ["GS ID", "Source", "Link", "Comment text", "1. Relevant?", "2. Main thing that happened to the SIP",
        "3. Drop-off type", "4. Main reason", "Notes (optional)"]
widths = [7, 13, 12, 80, 12, 30, 34, 32, 30]
for j, (h, w) in enumerate(zip(cols, widths), 1):
    c = S.cell(row=1, column=j, value=h); c.font = hdr; c.fill = fill_h; c.alignment = wrap; c.border = box
    S.column_dimensions[c.column_letter].width = w
n = len(G)
for i, g in enumerate(G, 2):
    vals = [f"GS{i-1:02d}", g["source"].replace("_", " "), g.get("url") or "", g["text"]]
    for j, v in enumerate(vals, 1):
        c = S.cell(row=i, column=j, value=v); c.font = body; c.alignment = wrap; c.border = box
    if g.get("url"):
        S.cell(row=i, column=3).hyperlink = g["url"]; S.cell(row=i, column=3).value = "open"
        S.cell(row=i, column=3).font = Font(name=F, size=10, color="0563C1", underline="single")
    for j in range(5, 10):
        c = S.cell(row=i, column=j); c.fill = fill_in; c.border = box; c.font = body; c.alignment = wrap
    S.row_dimensions[i].height = min(220, 15 * (len(g["text"]) // 95 + 2))
last = n + 1
for rng, src in ((f"E2:E{last}", "=Lists!$H$2:$H$3"), (f"F2:F{last}", f"=Lists!$A$2:$A${len(EVENTS)+1}"),
                 (f"G2:G{last}", f"=Lists!$C$2:$C${len(DROPS)+1}"), (f"H2:H{last}", f"=Lists!$E$2:$E${len(DRIVERS)+1}")):
    dv = DataValidation(type="list", formula1=src, allow_blank=True); S.add_data_validation(dv); dv.add(rng)
S.freeze_panes = "E2"

# ---- AI labels (hidden, blind) ----
A = wb.create_sheet("AI labels")
acols = ["GS ID", "AI relevant", "AI events (all)", "AI drop-off type", "AI drivers (all)", "AI unmet need",
         "AI evidence", "Model", "Prompt", "Stratum", "Guardrail check on AI output"]
for j, h in enumerate(acols, 1):
    c = A.cell(row=1, column=j, value=h); c.font = hdr; c.fill = fill_h
gv = 0
for i, g in enumerate(G, 2):
    out_text = " ".join([g.get("unmet_need", ""), g.get("expectation_gap", ""), g.get("emergent_theme", "")])
    adv = bool(ADVICE.search(out_text)); fund = bool(FUND.search(out_text)); gv += (adv or fund)
    viol = "VIOLATION: gives advice" if adv else ("VIOLATION: names a fund/AMC in AI-written text" if fund else "pass")
    vals = [f"GS{i-1:02d}", "Yes" if g["relevant"] else "No", "|" + "|".join(g["event"] or ["none"]) + "|",
            g["dropoff_type"], "|" + "|".join(g["drivers"] or ["none"]) + "|", g.get("unmet_need", ""),
            g.get("evidence_span", ""), g.get("_model"), g.get("_prompt"), g["stratum"],
            viol]
    for j, v in enumerate(vals, 1):
        A.cell(row=i, column=j, value=v).font = body
for j, w in enumerate([7, 10, 40, 16, 50, 40, 50, 22, 18, 28, 30], 1):
    A.column_dimensions[A.cell(row=1, column=j).column_letter].width = w
A.sheet_state = "hidden"

# ---- Agreement (formulas) ----
Ag = wb.create_sheet("Agreement", 2)
gcols = ["GS ID", "Labelled?", "Relevant match", "Event match", "Drop-off match", "Reason match"]
for j, h in enumerate(gcols, 1):
    c = Ag.cell(row=1, column=j, value=h); c.font = hdr; c.fill = fill_h; Ag.column_dimensions[c.column_letter].width = 16
nE, nD, nR = len(EVENTS) + 1, len(DROPS) + 1, len(DRIVERS) + 1
for i in range(2, last + 1):
    lr = f"'Label here'!"
    Ag.cell(row=i, column=1, value=f"={lr}A{i}")
    Ag.cell(row=i, column=2, value=f'=IF(COUNTA({lr}E{i}:{lr}H{i})=4,1,0)')
    Ag.cell(row=i, column=3, value=f'=IF(B{i}=1,IF({lr}E{i}=\'AI labels\'!B{i},1,0),"")')
    Ag.cell(row=i, column=4, value=f'=IF(B{i}=1,IF(ISNUMBER(SEARCH("|"&INDEX(Lists!$B$2:$B${nE},MATCH({lr}F{i},Lists!$A$2:$A${nE},0))&"|",\'AI labels\'!C{i})),1,0),"")')
    Ag.cell(row=i, column=5, value=f'=IF(B{i}=1,IF(INDEX(Lists!$D$2:$D${nD},MATCH({lr}G{i},Lists!$C$2:$C${nD},0))=\'AI labels\'!D{i},1,0),"")')
    Ag.cell(row=i, column=6, value=f'=IF(B{i}=1,IF(ISNUMBER(SEARCH("|"&INDEX(Lists!$F$2:$F${nR},MATCH({lr}H{i},Lists!$E$2:$E${nR},0))&"|",\'AI labels\'!E{i})),1,0),"")')
    for j in range(1, 7): Ag.cell(row=i, column=j).font = body
s = last + 2
Ag.cell(row=s, column=1, value="Items labelled").font = Font(name=F, bold=True)
Ag.cell(row=s, column=2, value=f"=SUM(B2:B{last})")
Ag.cell(row=s + 1, column=1, value="Agreement rate").font = Font(name=F, bold=True)
for j, col in enumerate("CDEF", 3):
    c = Ag.cell(row=s + 1, column=j, value=f'=IF($B${s}=0,"",AVERAGE({col}2:{col}{last}))'); c.number_format = "0.0%"
Ag.cell(row=s + 2, column=1, value="Overall (all 4 fields)").font = Font(name=F, bold=True)
c = Ag.cell(row=s + 2, column=3, value=f'=IF($B${s}=0,"",SUM(C2:F{last})/(4*$B${s}))'); c.number_format = "0.0%"
Ag.cell(row=s + 4, column=1, value="Guardrail violations in AI output (checked automatically: fund names or advice wording)").font = Font(name=F, bold=True)
Ag.cell(row=s + 4, column=6, value=f'=COUNTIF(\'AI labels\'!K2:K{last},"VIOLATION*")')
Ag.cell(row=s + 5, column=1, value="Event and Reason count as a match if your choice is among the AI's tags (the AI can tag more than one).").font = Font(name=F, italic=True, size=9)
Ag.column_dimensions["A"].width = 22
wb.save("/home/claude/engine/out/golden_set_label_me.xlsx")
print("saved; guardrail violations precomputed:", gv)
